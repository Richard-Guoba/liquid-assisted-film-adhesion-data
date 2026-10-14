from __future__ import annotations

import argparse
import sys
import tempfile
from dataclasses import asdict, dataclass
from numbers import Real
from pathlib import Path
from typing import Final

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT: Final = Path(__file__).resolve().parents[1]
DEFAULT_DATA_PATH: Final = PROJECT_ROOT / "data" / "force" / "force_data.xlsx"
DEFAULT_CONFIG_PATH: Final = PROJECT_ROOT / "config" / "force_config.xlsx"
DEFAULT_RESULT_PATH: Final = PROJECT_ROOT / "outputs" / "force" / "force_fit_results.xlsx"

CONFIG_COLUMNS: Final = (
    "test_id", "repeat_id", "substrate", "liquid", "thickness_um",
    "ideal_dry_time_h", "real_dry_time_h", "velocity_mm_min",
    "fit_curve_region_start_mm", "fit_curve_region_end_mm",
)
RESULT_COLUMNS: Final = (
    *CONFIG_COLUMNS,
    "actual_fit_region_start_mm", "actual_fit_region_end_mm",
    "slope_n_per_mm", "fit_intercept_n", "r_squared",
    "fit_point_count", "fit_plot",
)
RAW_COLUMNS: Final = ("Time (s)", "Load (N)", "Displacement (mm)")


@dataclass(frozen=True)
class ForceConfig:
    test_id: str
    repeat_id: int
    substrate: str
    liquid: str
    thickness_um: float
    ideal_dry_time_h: float
    real_dry_time_h: float | None
    velocity_mm_min: float
    fit_curve_region_start_mm: float
    fit_curve_region_end_mm: float


@dataclass(frozen=True)
class FitResult:
    test_id: str
    repeat_id: int
    substrate: str
    liquid: str
    thickness_um: float
    ideal_dry_time_h: float
    real_dry_time_h: float | None
    velocity_mm_min: float
    fit_curve_region_start_mm: float
    fit_curve_region_end_mm: float
    actual_fit_region_start_mm: float
    actual_fit_region_end_mm: float
    slope_n_per_mm: float
    fit_intercept_n: float
    r_squared: float
    fit_point_count: int
    fit_plot: str


def validate_force_curve(test_id: str, frame: pd.DataFrame) -> tuple[np.ndarray, np.ndarray]:
    missing = [column for column in RAW_COLUMNS if column not in frame.columns]
    if missing:
        raise ValueError(f"{test_id}: raw sheet is missing columns {missing}")

    for column in ("Displacement (mm)", "Load (N)"):
        for position, value in enumerate(frame[column], start=2):
            if isinstance(value, (bool, np.bool_)) or not isinstance(value, Real):
                raise ValueError(
                    f"{test_id}: {column} must contain only numeric values; "
                    f"invalid value at workbook row {position}"
                )

    displacement = frame["Displacement (mm)"].to_numpy(dtype=float)
    load = frame["Load (N)"].to_numpy(dtype=float)
    finite = np.isfinite(displacement) & np.isfinite(load)
    return displacement[finite], load[finite]


def fit_force_curve(config: ForceConfig, frame: pd.DataFrame) -> FitResult:
    displacement, load = validate_force_curve(config.test_id, frame)
    selected = (
        (displacement >= config.fit_curve_region_start_mm)
        & (displacement <= config.fit_curve_region_end_mm)
    )
    x_fit = displacement[selected]
    y_fit = load[selected]
    if x_fit.size < 2:
        raise ValueError(f"{config.test_id}: fit interval contains fewer than two finite points")
    if np.ptp(x_fit) == 0:
        raise ValueError(f"{config.test_id}: fit displacement has zero variance")
    if np.ptp(y_fit) == 0:
        raise ValueError(f"{config.test_id}: fit response has zero variance")

    slope, intercept = np.polyfit(x_fit, y_fit, 1)
    predicted = slope * x_fit + intercept
    residual_sum = float(np.sum((y_fit - predicted) ** 2))
    total_sum = float(np.sum((y_fit - np.mean(y_fit)) ** 2))
    r_squared = 1.0 - residual_sum / total_sum
    return FitResult(
        **asdict(config),
        actual_fit_region_start_mm=float(np.min(x_fit)),
        actual_fit_region_end_mm=float(np.max(x_fit)),
        slope_n_per_mm=float(slope),
        fit_intercept_n=float(intercept),
        r_squared=float(r_squared),
        fit_point_count=int(x_fit.size),
        fit_plot=f"fit_plots/{config.test_id}.png",
    )


def _finite_number(value: object, test_id: str, column: str) -> float:
    if isinstance(value, (bool, np.bool_)) or not isinstance(value, Real):
        raise ValueError(f"{test_id}: {column} must be numeric")
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{test_id}: {column} must be numeric") from error
    if not np.isfinite(number):
        raise ValueError(f"{test_id}: {column} must be finite")
    return number


def reject_workbook_formulas(path: Path, workbook_context: str) -> None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        raise ValueError(
                            f"{workbook_context} workbook {path}: sheet {sheet.title!r}, "
                            f"cell {cell.coordinate} contains a formula"
                        )
    finally:
        workbook.close()


def load_force_config(path: Path) -> list[ForceConfig]:
    reject_workbook_formulas(path, "force config")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if workbook.sheetnames != ["force_config"]:
            raise ValueError("force config workbook must contain only sheet 'force_config'")
        sheet = workbook["force_config"]
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers != CONFIG_COLUMNS:
            raise ValueError(f"force config columns must be exactly {list(CONFIG_COLUMNS)}")
        records = [
            dict(zip(CONFIG_COLUMNS, values, strict=True))
            for values in sheet.iter_rows(min_row=2, values_only=True)
        ]
    finally:
        workbook.close()
    ids: list[str] = []
    for raw in records:
        value = raw["test_id"]
        if not isinstance(value, str):
            raise ValueError("force config test_id must be text")
        test_id = value.strip()
        if not test_id:
            raise ValueError("force config contains an empty test_id")
        ids.append(test_id)
    duplicates = sorted({test_id for test_id in ids if ids.count(test_id) > 1})
    if duplicates:
        raise ValueError(f"force config contains duplicate test_id values: {duplicates}")

    configs: list[ForceConfig] = []
    for raw, test_id in zip(records, ids, strict=True):
        for column in ("substrate", "liquid"):
            value = raw[column]
            if not isinstance(value, str):
                raise ValueError(f"{test_id}: {column} must be text")
            value = value.strip()
            if not value:
                raise ValueError(f"{test_id}: {column} is required")
            raw[column] = value
        repeat_value = _finite_number(raw["repeat_id"], test_id, "repeat_id")
        if not repeat_value.is_integer():
            raise ValueError(f"{test_id}: repeat_id must be an integer")
        real_dry = None if pd.isna(raw["real_dry_time_h"]) else _finite_number(
            raw["real_dry_time_h"], test_id, "real_dry_time_h"
        )
        config = ForceConfig(
            test_id=test_id,
            repeat_id=int(repeat_value),
            substrate=raw["substrate"],
            liquid=raw["liquid"],
            thickness_um=_finite_number(raw["thickness_um"], test_id, "thickness_um"),
            ideal_dry_time_h=_finite_number(raw["ideal_dry_time_h"], test_id, "ideal_dry_time_h"),
            real_dry_time_h=real_dry,
            velocity_mm_min=_finite_number(raw["velocity_mm_min"], test_id, "velocity_mm_min"),
            fit_curve_region_start_mm=_finite_number(
                raw["fit_curve_region_start_mm"], test_id, "fit_curve_region_start_mm"
            ),
            fit_curve_region_end_mm=_finite_number(
                raw["fit_curve_region_end_mm"], test_id, "fit_curve_region_end_mm"
            ),
        )
        if config.fit_curve_region_start_mm > config.fit_curve_region_end_mm:
            raise ValueError(f"{test_id}: fit interval start exceeds end")
        configs.append(config)
    return configs


def natural_test_key(test_id: str) -> int:
    prefix, separator, number = test_id.rpartition("_")
    if prefix != "test" or separator != "_" or not number.isdigit():
        raise ValueError(f"invalid test_id format: {test_id}")
    return int(number)


def compute_force_results(data_path: Path, config_path: Path) -> list[FitResult]:
    configs = load_force_config(config_path)
    raw_book = pd.ExcelFile(data_path)
    configured_ids = {config.test_id for config in configs}
    raw_ids = set(raw_book.sheet_names)
    if configured_ids != raw_ids:
        missing = sorted(configured_ids - raw_ids)
        extra = sorted(raw_ids - configured_ids)
        raise ValueError(f"raw/config test_id mismatch; missing raw sheets={missing}; extra raw sheets={extra}")
    results = []
    for config in sorted(configs, key=lambda item: natural_test_key(item.test_id)):
        frame = pd.read_excel(raw_book, sheet_name=config.test_id)
        results.append(fit_force_curve(config, frame))
    return results


def results_frame(results: list[FitResult]) -> pd.DataFrame:
    frame = pd.DataFrame([asdict(result) for result in results], columns=RESULT_COLUMNS)
    frame["real_dry_time_h"] = frame["real_dry_time_h"].where(
        frame["real_dry_time_h"].notna(), np.nan
    )
    return frame


def write_results_atomic(results: list[FitResult], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output_path.parent,
            prefix=f".{output_path.stem}.",
            suffix=".tmp.xlsx",
            delete=False,
        ) as handle:
            temporary_path = Path(handle.name)
        frame = results_frame(results)
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            frame.to_excel(writer, sheet_name="force_fit_results", index=False)
            sheet = writer.book["force_fit_results"]
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.border = Border()
                cell.alignment = Alignment()
                cell.fill = PatternFill()
                cell.number_format = "General"
            for column_index, column_name in enumerate(RESULT_COLUMNS, start=1):
                values = [column_name, *frame[column_name].fillna("").astype(str).tolist()]
                sheet.column_dimensions[get_column_letter(column_index)].width = min(
                    max(len(value) for value in values) + 2,
                    32,
                )
        reject_workbook_formulas(temporary_path, "temporary force result")
        with pd.ExcelFile(temporary_path) as check:
            if check.sheet_names != ["force_fit_results"]:
                raise OSError("temporary result workbook has an invalid sheet layout")
            written = pd.read_excel(check, sheet_name="force_fit_results")
        if written.columns.tolist() != list(RESULT_COLUMNS) or len(written) != len(results):
            raise OSError("temporary result workbook failed row or column verification")
        temporary_path.replace(output_path)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def run_processing(data_path: Path, config_path: Path, output_path: Path) -> None:
    results = compute_force_results(data_path, config_path)
    if len(results) != len(load_force_config(config_path)):
        raise ValueError("force result/configuration counts differ")
    write_results_atomic(results, output_path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fit force-displacement curves from the project workbooks.")
    parser.add_argument("--data", type=Path, default=DEFAULT_DATA_PATH)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_RESULT_PATH)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        run_processing(args.data, args.config, args.output)
    except Exception as error:
        print(f"force processing failed: {error}", file=sys.stderr)
        return 1
    print(f"wrote force fit results: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
