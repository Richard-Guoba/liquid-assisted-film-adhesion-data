"""Strict loading and numerical fitting for the retained tensile-modulus curves."""

from __future__ import annotations

import argparse
import sys
import tempfile
from dataclasses import dataclass
from numbers import Real
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Final

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT: Final = Path(__file__).resolve().parents[1]
DEFAULT_DATA_PATH: Final = PROJECT_ROOT / "data" / "modulus" / "modulus_data.xlsx"
DEFAULT_CONFIG_PATH: Final = PROJECT_ROOT / "config" / "modulus_config.xlsx"
DEFAULT_RESULT_PATH: Final = PROJECT_ROOT / "outputs" / "modulus" / "modulus_results.xlsx"

CONFIG_COLUMNS: Final = (
    "test_id", "specimen_type", "gauge_length_mm", "width_mm", "thickness_um",
    "engineering_fit_start", "engineering_fit_end", "true_fit_start", "true_fit_end",
    "initial_true_fit_start", "initial_true_fit_end",
)
RESULT_COLUMNS: Final = (
    "test_id", "specimen_type", "gauge_length_mm", "width_mm", "thickness_um",
    "displacement_zero_mm", "status",
    "engineering_fit_start", "engineering_fit_end", "engineering_actual_start",
    "engineering_actual_end", "engineering_modulus_mpa", "engineering_intercept_mpa",
    "engineering_r_squared", "engineering_point_count",
    "true_fit_start", "true_fit_end", "true_actual_start", "true_actual_end",
    "true_modulus_mpa", "true_intercept_mpa", "true_r_squared", "true_point_count",
    "initial_true_fit_start", "initial_true_fit_end", "initial_true_actual_start",
    "initial_true_actual_end", "initial_true_modulus_mpa", "initial_true_intercept_mpa",
    "initial_true_r_squared", "initial_true_point_count", "fit_plot",
)
RAW_COLUMNS: Final = ("Force (N)", "Displacement (mm)")
EXPECTED_IDS: Final = tuple(f"test_{index:02d}" for index in range(1, 54))
_INTERVAL_COLUMNS: Final = (
    ("engineering_fit_start", "engineering_fit_end"),
    ("true_fit_start", "true_fit_end"),
    ("initial_true_fit_start", "initial_true_fit_end"),
)


@dataclass(frozen=True)
class ModulusConfig:
    test_id: str
    specimen_type: str
    gauge_length_mm: float
    width_mm: float
    thickness_um: float
    engineering_fit_start: float
    engineering_fit_end: float
    true_fit_start: float
    true_fit_end: float
    initial_true_fit_start: float
    initial_true_fit_end: float


@dataclass(frozen=True)
class LinearFit:
    start: float
    end: float
    actual_start: float
    actual_end: float
    slope_mpa: float
    intercept_mpa: float
    r_squared: float
    point_count: int


@dataclass(frozen=True)
class CurveData:
    displacement_zero_mm: float
    engineering_strain: np.ndarray
    engineering_stress_mpa: np.ndarray
    true_strain: np.ndarray
    true_stress_mpa: np.ndarray


@dataclass(frozen=True)
class ModulusResult:
    config: ModulusConfig
    displacement_zero_mm: float
    engineering: LinearFit | None
    true: LinearFit | None
    initial_true: LinearFit | None
    status: str
    fit_plot: str


class _InsufficientFitError(ValueError):
    """A fit cannot be calculated from the available, otherwise valid curve data."""


def _reject_workbook_formulas(path: Path) -> None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        raise ValueError(
                            f"{sheet.title}: cell {cell.coordinate} contains a formula"
                        )
    finally:
        workbook.close()


def _finite_number(value: object, test_id: str, column: str) -> float:
    if isinstance(value, (bool, np.bool_)) or not isinstance(value, Real):
        raise ValueError(f"{test_id}: {column} must be a numeric value")
    number = float(value)
    if not np.isfinite(number):
        raise ValueError(f"{test_id}: {column} must be finite")
    return number


def _validate_config_record(record: dict[str, object]) -> ModulusConfig:
    test_id = record["test_id"]
    if not isinstance(test_id, str) or not test_id:
        raise ValueError("modulus_config: test_id must be nonempty text")
    specimen_type = record["specimen_type"]
    if not isinstance(specimen_type, str) or not specimen_type:
        raise ValueError(f"{test_id}: specimen_type must be nonempty text")

    geometry = {
        column: _finite_number(record[column], test_id, column)
        for column in ("gauge_length_mm", "width_mm", "thickness_um")
    }
    for column, value in geometry.items():
        if value <= 0.0:
            raise ValueError(f"{test_id}: {column} must be positive")
    intervals = {
        column: _finite_number(record[column], test_id, column)
        for pair in _INTERVAL_COLUMNS for column in pair
    }
    for start_column, end_column in _INTERVAL_COLUMNS:
        start, end = intervals[start_column], intervals[end_column]
        if start < 0.0 or start >= end:
            raise ValueError(f"{test_id}: {start_column}/{end_column} must satisfy 0 <= start < end")
    return ModulusConfig(test_id=test_id, specimen_type=specimen_type, **geometry, **intervals)


def load_modulus_config(path: Path) -> list[ModulusConfig]:
    """Load the exact, formula-free 53-row modulus configuration workbook."""
    _reject_workbook_formulas(path)
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if workbook.sheetnames != ["modulus_config"]:
            raise ValueError("modulus config workbook must contain only sheet 'modulus_config'")
        sheet = workbook["modulus_config"]
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers != CONFIG_COLUMNS:
            raise ValueError(f"modulus_config columns must be exactly {list(CONFIG_COLUMNS)}")
        records = [
            dict(zip(CONFIG_COLUMNS, values, strict=True))
            for values in sheet.iter_rows(min_row=2, values_only=True)
        ]
    finally:
        workbook.close()

    configs = [_validate_config_record(record) for record in records]
    ids = tuple(config.test_id for config in configs)
    if len(set(ids)) != len(ids):
        raise ValueError("modulus_config contains duplicate test_id values")
    if ids != EXPECTED_IDS:
        raise ValueError("modulus_config test_id values must be test_01 through test_53 in order")
    return configs


def _is_missing(value: object) -> bool:
    if value is None:
        return True
    missing = pd.isna(value)
    return isinstance(missing, (bool, np.bool_)) and bool(missing)


def validate_raw_curve(test_id: str, frame: pd.DataFrame) -> tuple[np.ndarray, np.ndarray]:
    """Reject malformed values and retain rows with finite force and displacement."""
    if tuple(frame.columns) != RAW_COLUMNS:
        raise ValueError(f"{test_id}: raw sheet columns must be exactly {list(RAW_COLUMNS)}")
    for column in RAW_COLUMNS:
        for row_number, value in enumerate(frame[column], start=2):
            if _is_missing(value):
                continue
            if isinstance(value, (bool, np.bool_)) or not isinstance(value, Real):
                raise ValueError(
                    f"{test_id}: {column} must contain only numeric values; invalid value at row {row_number}"
                )
    force = pd.to_numeric(frame["Force (N)"], errors="raise").to_numpy(dtype=float)
    displacement = pd.to_numeric(frame["Displacement (mm)"], errors="raise").to_numpy(dtype=float)
    retained = np.isfinite(force) & np.isfinite(displacement)
    return force[retained], displacement[retained]


def convert_curve(config: ModulusConfig, frame: pd.DataFrame) -> CurveData:
    """Convert one raw force/displacement table using the documented strain equations."""
    # Validate types before deciding which rows are usable; the reference displacement
    # is taken from the first finite displacement, even if its force is not finite.
    if tuple(frame.columns) != RAW_COLUMNS:
        raise ValueError(f"{config.test_id}: raw sheet columns must be exactly {list(RAW_COLUMNS)}")
    for column in RAW_COLUMNS:
        for row_number, value in enumerate(frame[column], start=2):
            if _is_missing(value):
                continue
            if isinstance(value, (bool, np.bool_)) or not isinstance(value, Real):
                raise ValueError(
                    f"{config.test_id}: {column} must contain only numeric values; invalid value at row {row_number}"
                )
    force_all = pd.to_numeric(frame["Force (N)"], errors="raise").to_numpy(dtype=float)
    displacement_all = pd.to_numeric(
        frame["Displacement (mm)"], errors="raise"
    ).to_numpy(dtype=float)
    finite_displacement = np.isfinite(displacement_all)
    if not np.any(finite_displacement):
        raise ValueError(f"{config.test_id}: raw sheet has no finite displacement")
    displacement_zero = float(displacement_all[finite_displacement][0])
    retained = np.isfinite(force_all) & finite_displacement
    force = force_all[retained]
    displacement = displacement_all[retained]

    extension = displacement - displacement_zero
    engineering_strain = extension / config.gauge_length_mm
    engineering_stress = force / (config.width_mm * config.thickness_um / 1000.0)
    true_valid = 1.0 + engineering_strain > 0.0
    return CurveData(
        displacement_zero_mm=displacement_zero,
        engineering_strain=engineering_strain,
        engineering_stress_mpa=engineering_stress,
        true_strain=np.log1p(engineering_strain[true_valid]),
        true_stress_mpa=engineering_stress[true_valid] * (1.0 + engineering_strain[true_valid]),
    )


def fit_interval(
    x: np.ndarray, y: np.ndarray, start: float, end: float, test_id: str
) -> LinearFit:
    """Fit a straight line by inclusive ordinary least squares."""
    selected = np.isfinite(x) & np.isfinite(y) & (x >= start) & (x <= end)
    x_fit, y_fit = x[selected], y[selected]
    if x_fit.size < 2:
        raise _InsufficientFitError(f"{test_id}: fit interval contains fewer than two finite points")
    if np.ptp(x_fit) == 0.0:
        raise _InsufficientFitError(f"{test_id}: fit strain has zero variance")
    slope, intercept = np.polyfit(x_fit, y_fit, 1)
    predicted = slope * x_fit + intercept
    ss_res = float(np.sum((y_fit - predicted) ** 2))
    ss_tot = float(np.sum((y_fit - np.mean(y_fit)) ** 2))
    r_squared = 1.0 - ss_res / ss_tot if ss_tot > 0.0 else 1.0
    return LinearFit(
        start=float(start), end=float(end), actual_start=float(np.min(x_fit)),
        actual_end=float(np.max(x_fit)), slope_mpa=float(slope),
        intercept_mpa=float(intercept), r_squared=float(r_squared), point_count=int(x_fit.size),
    )


def _fit_or_unavailable(
    x: np.ndarray, y: np.ndarray, start: float, end: float, test_id: str
) -> LinearFit | None:
    try:
        return fit_interval(x, y, start, end, test_id)
    except _InsufficientFitError:
        return None


def _validate_raw_workbook(path: Path) -> None:
    _reject_workbook_formulas(path)
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if tuple(workbook.sheetnames) != EXPECTED_IDS:
            raise ValueError("modulus raw workbook sheets must be test_01 through test_53 in order")
        for test_id in EXPECTED_IDS:
            sheet = workbook[test_id]
            headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
            if headers != RAW_COLUMNS:
                raise ValueError(f"{test_id}: raw sheet columns must be exactly {list(RAW_COLUMNS)}")
    finally:
        workbook.close()


def _load_raw_frames(path: Path) -> dict[str, pd.DataFrame]:
    """Read raw cells without pandas' Excel type inference."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        return {
            test_id: pd.DataFrame(
                list(workbook[test_id].iter_rows(min_row=2, values_only=True)),
                columns=RAW_COLUMNS,
            )
            for test_id in EXPECTED_IDS
        }
    finally:
        workbook.close()


def compute_modulus_results(data_path: Path, configs: list[ModulusConfig]) -> list[ModulusResult]:
    """Compute all three configured fits while preserving tests with unavailable fits."""
    if tuple(config.test_id for config in configs) != EXPECTED_IDS:
        raise ValueError("configs must be test_01 through test_53 in order")
    _validate_raw_workbook(data_path)
    frames = _load_raw_frames(data_path)
    results: list[ModulusResult] = []
    for config in configs:
        curve = convert_curve(config, frames[config.test_id])
        engineering = _fit_or_unavailable(
            curve.engineering_strain, curve.engineering_stress_mpa,
            config.engineering_fit_start, config.engineering_fit_end, config.test_id,
        )
        true = _fit_or_unavailable(
            curve.true_strain, curve.true_stress_mpa,
            config.true_fit_start, config.true_fit_end, config.test_id,
        )
        initial_true = _fit_or_unavailable(
            curve.true_strain, curve.true_stress_mpa,
            config.initial_true_fit_start, config.initial_true_fit_end, config.test_id,
        )
        unavailable = [
            name for name, fit in (
                ("engineering_fit_unavailable", engineering),
                ("true_fit_unavailable", true),
                ("initial_true_fit_unavailable", initial_true),
            ) if fit is None
        ]
        results.append(ModulusResult(
            config=config, displacement_zero_mm=curve.displacement_zero_mm,
            engineering=engineering, true=true, initial_true=initial_true,
            status=";".join(unavailable) if unavailable else "ok",
            fit_plot=f"fit_plots/{config.test_id}.png",
        ))
    return results


def _natural_test_key(test_id: str) -> int:
    prefix, separator, number = test_id.rpartition("_")
    if prefix != "test" or separator != "_" or not number.isdigit():
        raise ValueError(f"invalid test_id format: {test_id}")
    return int(number)


def _fit_values(prefix: str, fit: LinearFit | None) -> dict[str, float | int | None]:
    fields = (
        ("fit_start", "start"), ("fit_end", "end"),
        ("actual_start", "actual_start"), ("actual_end", "actual_end"),
        ("modulus_mpa", "slope_mpa"), ("intercept_mpa", "intercept_mpa"),
        ("r_squared", "r_squared"), ("point_count", "point_count"),
    )
    return {
        f"{prefix}_{result_name}": None if fit is None else getattr(fit, attribute)
        for result_name, attribute in fields
    }


def _result_row(result: ModulusResult) -> dict[str, object]:
    windows_path = PureWindowsPath(result.fit_plot) if isinstance(result.fit_plot, str) else None
    if (
        not isinstance(result.fit_plot, str)
        or not result.fit_plot
        or PurePosixPath(result.fit_plot).is_absolute()
        or windows_path.drive
        or windows_path.root
    ):
        raise ValueError(f"{result.config.test_id}: fit_plot must be a nonempty relative path")
    return {
        "test_id": result.config.test_id,
        "specimen_type": result.config.specimen_type,
        "gauge_length_mm": result.config.gauge_length_mm,
        "width_mm": result.config.width_mm,
        "thickness_um": result.config.thickness_um,
        "displacement_zero_mm": result.displacement_zero_mm,
        "status": result.status,
        **_fit_values("engineering", result.engineering),
        **_fit_values("true", result.true),
        **_fit_values("initial_true", result.initial_true),
        "fit_plot": result.fit_plot,
    }


def results_frame(results: list[ModulusResult]) -> pd.DataFrame:
    """Flatten fit results into the fixed, naturally ordered publication schema."""
    rows = [_result_row(result) for result in results]
    frame = pd.DataFrame(rows, columns=RESULT_COLUMNS)
    if not frame.empty:
        frame = frame.sort_values(
            "test_id", key=lambda series: series.map(_natural_test_key), kind="stable"
        ).reset_index(drop=True)
    return frame


def _validate_temporary_results(path: Path, expected_rows: int) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        if workbook.sheetnames != ["modulus_results"]:
            raise OSError("temporary result workbook has an invalid sheet layout")
        sheet = workbook["modulus_results"]
        if sheet.max_row != expected_rows + 1 or sheet.max_column != len(RESULT_COLUMNS):
            raise OSError("temporary result workbook failed row or column verification")
        if tuple(cell.value for cell in sheet[1]) != RESULT_COLUMNS:
            raise OSError("temporary result workbook has invalid result columns")
        if sheet.auto_filter.ref is not None or sheet.tables or sheet.merged_cells.ranges or sheet._charts:
            raise OSError("temporary result workbook has unsupported decorations")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f" or cell.comment is not None:
                    raise OSError("temporary result workbook contains formulas or comments")
    finally:
        workbook.close()


def write_results_atomic(results: list[ModulusResult], output_path: Path) -> None:
    """Write a structurally validated result workbook, replacing only on success."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output_path.parent,
            prefix=f".{output_path.stem}.",
            suffix=".tmp.xlsx",
            delete=False,
        ) as handle:
            temporary_path = Path(handle.name)
        frame = results_frame(results)
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            frame.to_excel(writer, sheet_name="modulus_results", index=False)
            sheet = writer.book["modulus_results"]
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.border = Border()
                cell.alignment = Alignment()
                cell.fill = PatternFill()
                cell.number_format = "General"
            for column_index, column_name in enumerate(RESULT_COLUMNS, start=1):
                values = [column_name, *frame[column_name].fillna("").astype(str).tolist()]
                sheet.column_dimensions[get_column_letter(column_index)].width = min(
                    max(len(value) for value in values) + 2, 32
                )
        _validate_temporary_results(temporary_path, len(frame))
        temporary_path.replace(output_path)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def run_processing(data_path: Path, config_path: Path, output_path: Path) -> None:
    results = compute_modulus_results(data_path, load_modulus_config(config_path))
    if data_path == DEFAULT_DATA_PATH and config_path == DEFAULT_CONFIG_PATH and len(results) != 53:
        raise ValueError(f"default modulus run must produce 53 results, received {len(results)}")
    write_results_atomic(results, output_path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fit tensile-modulus curves from the project workbooks.")
    parser.add_argument("--data", type=Path, default=DEFAULT_DATA_PATH)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_RESULT_PATH)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        run_processing(args.data, args.config, args.output)
    except Exception as error:
        print(f"modulus processing failed: {error}", file=sys.stderr)
        return 1
    print(f"wrote modulus fit results: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
